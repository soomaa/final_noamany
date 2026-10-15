from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "screenshots"
OUT = ROOT / "ONE80-Cafe-Verification-Report-2026-07-18.docx"

BLUE = "2E74B5"
DARK = "1F4D78"
LIGHT = "F2F4F7"
GREEN = "1B7F5A"
MUTED = "666666"
GOLD = "C9A227"


def set_run(run, size=11, bold=False, color="000000", italic=False):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def rtl_paragraph(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    pPr = p._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
    return p


def add_rtl(doc, text, size=11, bold=False, color="000000", after=6, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = doc.add_paragraph()
    rtl_paragraph(p, align)
    p.paragraph_format.space_after = Pt(after)
    set_run(p.add_run(text), size=size, bold=bold, color=color)
    return p


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(width_dxa))
    tcW.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "120")
    tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = tcPr.find(qn("w:tcMar"))
            if tcMar is None:
                tcMar = OxmlElement("w:tcMar")
                tcPr.append(tcMar)
            for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                node = tcMar.find(qn(f"w:{side}"))
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    tcMar.append(node)
                node.set(qn("w:w"), str(value))
                node.set(qn("w:type"), "dxa")


def cell_text(cell, text, bold=False, color="000000", center=False, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    rtl_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.RIGHT)
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(str(text)), size=size, bold=bold, color=color)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    rtl_paragraph(p)
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    run = p.add_run(text)
    set_run(run, size=16 if level == 1 else 13, bold=True, color=BLUE if level < 3 else DARK)
    return p


def add_status_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_cell_shading(table.cell(0, 0), "EAF4EF")
    p = table.cell(0, 0).paragraphs[0]
    rtl_paragraph(p)
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(title + " — "), size=11, bold=True, color=GREEN)
    set_run(p.add_run(text), size=10.5, color="20352D")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc, filename, caption):
    path = SHOTS / filename
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(6.2))
    cap = doc.add_paragraph()
    rtl_paragraph(cap, WD_ALIGN_PARAGRAPH.CENTER)
    cap.paragraph_format.space_after = Pt(10)
    set_run(cap.add_run(caption), size=9, italic=True, color=MUTED)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
for name, size, color, before, after in (("Heading 1", 16, BLUE, 16, 8), ("Heading 2", 13, BLUE, 12, 6), ("Heading 3", 12, DARK, 8, 4)):
    st = styles[name]
    st.font.name = "Arial"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

header = sec.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run(hp.add_run("ONE80 | تقرير التحقق من نظام الكافيه والمخزون"), size=8.5, color=MUTED)
footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run(fp.add_run("ONE80 — 18 يوليو 2026"), size=8.5, color=MUTED)

# Memo masthead opening.
add_rtl(doc, "تقرير التحقق الفني", size=10, bold=True, color=GOLD, after=4)
add_rtl(doc, "نظام الكافيه والمخزون", size=24, bold=True, color=DARK, after=4)
add_rtl(doc, "مشروع ONE80 — تحقق وظيفي وربط مالي ومخزني", size=13, color=MUTED, after=16)

meta = doc.add_table(rows=4, cols=2)
meta_data = [("تاريخ الاختبار", "18 يوليو 2026"), ("قاعدة البيانات", "one80_hr المنفصلة"), ("نطاق الاختبار", "13 شاشة من دورة الكافيه"), ("حالة البناء", "Frontend وBackend ناجحان")]
for i, (label, value) in enumerate(meta_data):
    cell_text(meta.cell(i, 0), label, bold=True, color=DARK, size=10)
    cell_text(meta.cell(i, 1), value, size=10)
    set_cell_shading(meta.cell(i, 0), LIGHT)
set_table_geometry(meta, [2500, 6860])

add_heading(doc, "النتيجة التنفيذية", 1)
add_status_callout(doc, "النتيجة", "النظام مترابط ويعمل من المورد والمشتريات إلى المخزون والمديونية والجرد والحركات والتقارير.")
for label, value in [
    ("Migrations", "تم تطبيق 27 migration ولا توجد migrations معلّقة."),
    ("Production Build", "نجح بناء الواجهة الأمامية والـ backend."),
    ("سلامة البيانات", "لم يتم لمس بيانات الأعضاء أو الاشتراكات، ولم تتغير هوية ONE80."),
    ("أخطاء التشغيل", "لم يظهر خطأ قاعدة بيانات، وسجل Console خالٍ من الأخطاء والتحذيرات أثناء الاختبار."),
    ("اختبار فعلي", "فاتورة شراء 100 ج.م ومدفوع 40 ج.م رفعت المخزون وسجلت مديونية 60 ج.م؛ ثم سداد 20 ج.م خفّض المتبقي إلى 40 ج.م."),
]:
    p = add_rtl(doc, "", after=4)
    set_run(p.add_run(label + ": "), size=11, bold=True, color=DARK)
    set_run(p.add_run(value), size=11)

doc.add_page_break()
add_heading(doc, "نتائج الصفحات", 1)
rows = [
    ("فئات الكافيه", "تنظيم تبويبات نقطة البيع والتقارير دون تصنيف الخامات", "ناجح"),
    ("خامات الكافيه", "تنبيه كمية، متوسط تكلفة، ووحدات oz وfl oz بدون سعر بيع", "ناجح"),
    ("منتجات الكافيه", "صورة ووصفة وتكلفة وربح وثلاثة أنواع للمنتجات", "ناجح"),
    ("قائمة الأسعار", "تعديل سعر المنتجات القابلة للبيع وانعكاسه في نقطة البيع", "ناجح"),
    ("الموردون", "إنشاء مورد وربطه بالخامات التي يوفرها", "ناجح"),
    ("المشتريات", "فلترة الخامات وحساب سعر الوحدة والمدفوع والمتبقي", "ناجح فعليًا"),
    ("مديونية الموردين", "تحديث المديونية وآخر سداد وتوزيع الدفعة على الأقدم", "ناجح فعليًا"),
    ("إيصال بيع جديد", "عميل/موظف/شريك وسلة ودفع بدون ورديات", "ناجح"),
    ("الجرد", "بداية اليوم والوارد والمنصرف والاستهلاك والإغلاق والموجود", "ناجح"),
    ("سجل الحركات", "المصدر والمستند والمستخدم والرصيد بعد الحركة", "ناجح"),
    ("لوحة المخزون", "القيمة والمنخفض والوارد والمنصرف والحركات والتحليلات", "ناجح"),
    ("صرف داخل الجيم", "تحويل مخزون بدون معاملة مالية", "ناجح"),
    ("تقارير الكافيه", "مبيعات وربح ومشتريات ومديونية وأنواع البيع", "ناجح"),
]
tbl = doc.add_table(rows=1, cols=3)
for i, h in enumerate(("الصفحة", "ما تم التحقق منه", "النتيجة")):
    cell_text(tbl.cell(0, i), h, bold=True, color="FFFFFF", center=True, size=10)
    set_cell_shading(tbl.cell(0, i), BLUE)
for page, check, status in rows:
    cells = tbl.add_row().cells
    cell_text(cells[0], page, bold=True, color=DARK, size=9)
    cell_text(cells[1], check, size=9)
    cell_text(cells[2], status, bold=True, color=GREEN, center=True, size=9)
set_repeat_table_header(tbl.rows[0])
set_table_geometry(tbl, [2100, 5760, 1500])

doc.add_page_break()
add_heading(doc, "التحقق من قواعد العمل", 1)
rules = [
    ("الخامات", "لا يوجد سعر بيع؛ متوسط التكلفة يتغير من المشتريات والحركات فقط."),
    ("تنبيه الكمية", "قابل للتعديل ويظهر عند وصول الرصيد للحد المحدد أو أقل."),
    ("الأونصة", "وحدتا oz وfl oz متاحتان في نموذج الخامة."),
    ("أنواع المنتجات", "محضّر للبيع، تصنيع داخلي، ومنتج جاهز؛ الداخلي لا يظهر في نقطة البيع."),
    ("الوصفة", "تقبل مكونًا موجودًا أو إنشاء خامة جديدة، مع حساب التكلفة والربح تلقائيًا."),
    ("المشتريات", "فلترة خامات المورد وحساب سعر الوحدة من قيمة الدفعة ÷ الكمية."),
    ("المخزون", "فاتورة الاختبار رفعت الرصيد من 0 إلى 10 أونصات وسجلت حركة."),
    ("المديونية", "100 إجمالي، 40 مدفوع، 60 متبقي؛ وبعد سداد 20 أصبح المتبقي 40."),
    ("الوردية", "لا يوجد اعتماد على الوردية؛ البيع مسجل باسم المستخدم الحالي."),
    ("التقارير", "الشراء انعكس في لوحة المخزون والجرد والحركات وتقارير الكافيه."),
]
for label, text in rules:
    p = add_rtl(doc, "", after=5)
    set_run(p.add_run(label + ": "), size=11, bold=True, color=DARK)
    set_run(p.add_run(text), size=11)

add_heading(doc, "بيانات الاختبار المحلية", 2)
add_rtl(doc, "جميع السجلات المستخدمة تحمل البادئة E2E-CAFE حتى يمكن تمييزها وحذفها من قاعدة ONE80 المحلية فقط.")
data_tbl = doc.add_table(rows=1, cols=2)
for i, h in enumerate(("العنصر", "القيمة")):
    cell_text(data_tbl.cell(0, i), h, bold=True, color="FFFFFF", center=True)
    set_cell_shading(data_tbl.cell(0, i), BLUE)
for label, value in [("الفئة", "E2E-CAFE مشروبات اختبار"), ("الخامة", "E2E-CAFE كاكاو أونصة"), ("المورد", "E2E-CAFE مورد الاختبار"), ("المستودع", "E2E مخزن اختبار الكافيه"), ("فاتورة الشراء", "SINV-000001")]:
    row = data_tbl.add_row().cells
    cell_text(row[0], label, bold=True, color=DARK)
    cell_text(row[1], value)
set_table_geometry(data_tbl, [2500, 6860])

doc.add_page_break()
add_heading(doc, "أدلة الواجهات", 1)
figures = [
    ("01-cafe-categories.png", "الشكل 1 — فئات الكافيه"),
    ("02b-raw-material-form.png", "الشكل 2 — نموذج الخامة ووحدات الأونصة وتنبيه الكمية"),
    ("03b-product-form.png", "الشكل 3 — نموذج المنتج والصورة والوصفة والتكلفة"),
    ("06b-purchase-filtered-material.png", "الشكل 4 — فلترة خامات المورد في فاتورة الشراء"),
    ("06c-purchase-created.png", "الشكل 5 — فاتورة الشراء بعد الحفظ"),
    ("07c-supplier-payment-created.png", "الشكل 6 — المديونية بعد تسجيل دفعة المورد"),
    ("08-pos.png", "الشكل 7 — شاشة إيصال بيع جديد"),
    ("09b-stock-after-purchase.png", "الشكل 8 — متابعة مخزون اليوم بعد الشراء"),
    ("10b-movement-after-purchase.png", "الشكل 9 — سجل الحركة والمستخدم المنفذ"),
    ("11b-dashboard-after-purchase.png", "الشكل 10 — لوحة المخزون بعد الشراء"),
    ("12-gym-issue.png", "الشكل 11 — صرف مخزون داخل الجيم بدون معاملة مالية"),
    ("13b-reports-after-purchase.png", "الشكل 12 — تقارير الكافيه بعد الشراء"),
]
for idx, (filename, caption) in enumerate(figures):
    add_figure(doc, filename, caption)
    if idx in (2, 5, 8):
        doc.add_page_break()

doc.add_page_break()
add_heading(doc, "قائمة النشر الآمن", 1)
deploy = [
    ("Migration", "تشغيل npx prisma migrate deploy على قاعدة ONE80 قبل تشغيل النسخة الجديدة."),
    ("المستودعات", "التأكد من وجود مستودع رئيسي واحد على الأقل لكل فرع."),
    ("الصور", "حفظ uploads في persistent storage وعدم استبدالها أثناء النشر."),
    ("قاعدة البيانات", "عدم تشغيل seed أو reset على السيرفر."),
    ("النسخ الاحتياطي", "أخذ نسخة من قاعدة البيانات ومجلد الصور قبل أول نشر."),
]
for label, text in deploy:
    p = add_rtl(doc, "", after=6)
    set_run(p.add_run(label + ": "), size=11, bold=True, color=DARK)
    set_run(p.add_run(text), size=11)

add_status_callout(doc, "الحكم النهائي", "الدورة الأساسية للكافيه في ONE80 مترابطة، والصفحات المطلوبة تعمل دون أخطاء قاعدة بيانات، وعمليات الشراء والمخزون والمديونية والسداد انعكست في السجلات والتقارير كما هو مطلوب.")

doc.core_properties.title = "تقرير التحقق من نظام الكافيه والمخزون — ONE80"
doc.core_properties.subject = "Cafe and inventory verification report"
doc.core_properties.author = "ONE80"
doc.save(OUT)
print(OUT)
